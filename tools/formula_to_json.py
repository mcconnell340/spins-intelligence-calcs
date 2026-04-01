"""
formula_to_json.py
------------------
Converts a human-readable SPINS measure formula (e.g. "SUM(DOLLARS) / MAX(HOUSEHOLD)")
into the JSON step array used by Whiz.

Uses structural template matching against beep bop.xlsx — no API key required.

Usage:
    python tools/formula_to_json.py --name "My Measure" --formula "SUM(DOLLARS) / SUM(UNITS)"
    python tools/formula_to_json.py --batch input.csv --output output.csv
"""

import argparse
import copy
import csv
import json
import re
import sys
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
BEEP_BOP_PATH = ROOT / "beep bop.xlsx"


# ── Load examples ─────────────────────────────────────────────────────────────
def load_examples() -> list[dict]:
    wb = openpyxl.load_workbook(BEEP_BOP_PATH)
    ws = wb["beep bop"]
    examples = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, calc_json, calc_clean = row
        if name and calc_json and calc_clean and calc_clean.strip() != "-":
            try:
                parsed = json.loads(calc_json)
                examples.append({
                    "name": name,
                    "formula": calc_clean.strip(),
                    "json": parsed,
                })
            except json.JSONDecodeError:
                pass
    return examples


# ── Formula normalization ─────────────────────────────────────────────────────
def normalize_formula(formula: str):
    """
    Strip whitespace, uppercase, replace column names with C0/C1/... placeholders.
    Returns (normalized_string, ordered_list_of_(agg, col)_pairs).
    """
    formula = re.sub(r"\s+", "", formula.upper())
    seen: dict[str, str] = {}
    order: list[tuple[str, str]] = []
    counter = [0]

    def rep(m: re.Match) -> str:
        agg, col = m.group(1), m.group(2)
        if col not in seen:
            seen[col] = f"C{counter[0]}"
            order.append((agg, col))
            counter[0] += 1
        return f"{agg}({seen[col]})"

    normalized = re.sub(r"(SUM|MAX)\(([^)]+)\)", rep, formula)
    return normalized, order


# ── Template matching ─────────────────────────────────────────────────────────
def find_best_match(formula: str, examples: list[dict]) -> tuple[dict | None, float]:
    """
    Find the example whose structural signature best matches the input formula.
    Returns (example, score) where score=1.0 means exact structural match.
    """
    input_norm, _ = normalize_formula(formula)

    best_example = None
    best_score = 0.0

    for ex in examples:
        ex_norm, _ = normalize_formula(ex["formula"])
        if ex_norm == input_norm:
            return ex, 1.0
        score = SequenceMatcher(None, input_norm, ex_norm).ratio()
        if score > best_score:
            best_score = score
            best_example = ex

    return best_example, best_score


# ── Column substitution ───────────────────────────────────────────────────────
def apply_template(template_steps: list[dict], col_mapping: dict[str, str]) -> list[dict]:
    """
    Substitute column names (and dependent step names) into a copy of template_steps.
    col_mapping: {old_column_name: new_column_name}

    Strategy:
      Pass 1 — sum/max steps: always name = new_col.lower(). Builds name_map of
               old_step_name -> new_step_name for every aggregation step.
      Pass 2 — derived steps (in array order): update all reference fields via
               name_map, then derive new step name by replacing any old column
               names embedded in it (longest first to avoid partial collisions),
               then lowercase the result.
    """
    steps = copy.deepcopy(template_steps)
    name_map: dict[str, str] = {}

    # Sort old column names longest-first so longer names are replaced before
    # shorter substrings (e.g. DOLLARS_2YAGO before DOLLARS).
    col_replacements = sorted(col_mapping.items(), key=lambda kv: -len(kv[0]))

    # Pass 1: aggregation steps — name is always new_col.lower()
    for step in steps:
        if "column" not in step:
            continue
        old_col = step["column"]
        new_col = col_mapping.get(old_col, old_col)
        step["column"] = new_col
        old_name = step["name"]
        new_name = new_col.lower()
        step["name"] = new_name
        name_map[old_name] = new_name

    # Pass 2: derived steps — update references then fix step name
    for step in steps:
        if "column" in step:
            continue
        old_name = step["name"]

        # Update all reference fields
        for field in ("value", "denominator", "from", "base", "outOf"):
            if field in step:
                step[field] = name_map.get(step[field], step[field])
        if "values" in step:
            step["values"] = [name_map.get(v, v) for v in step["values"]]

        # Derive new step name: replace any embedded old column names, then lowercase
        new_name = old_name
        for old_col, new_col in col_replacements:
            new_name = re.sub(re.escape(old_col), new_col, new_name, flags=re.IGNORECASE)
        new_name = new_name.lower()
        step["name"] = new_name
        name_map[old_name] = new_name

    return steps


# ── Main conversion ───────────────────────────────────────────────────────────
def convert(name: str, formula: str, examples: list[dict]) -> tuple[list[dict], dict, float]:
    """
    Convert a clean formula to Whiz JSON steps using template matching.
    Returns (steps, matched_example, confidence_score).
    Raises ValueError if no match is found.
    """
    _, input_cols = normalize_formula(formula)
    best_ex, score = find_best_match(formula, examples)

    if best_ex is None:
        raise ValueError("No matching template found in beep bop.xlsx")

    _, template_cols = normalize_formula(best_ex["formula"])

    # Map template columns → input columns by position
    col_mapping = {
        tcol: input_cols[i][1]
        for i, (_, tcol) in enumerate(template_cols)
        if i < len(input_cols)
    }

    result = apply_template(best_ex["json"], col_mapping)
    return result, best_ex, score


# ── CLI ───────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Convert SPINS formula to Whiz JSON")
    parser.add_argument("--name", help="Measure name")
    parser.add_argument("--formula", help="Clean formula string")
    parser.add_argument("--batch", help="CSV with 'name' and 'formula' columns")
    parser.add_argument("--output", help="Output CSV path (batch mode)")
    args = parser.parse_args()

    examples = load_examples()
    print(f"Loaded {len(examples)} templates from beep bop.xlsx", file=sys.stderr)

    if args.batch:
        input_path = Path(args.batch)
        output_path = (
            Path(args.output) if args.output
            else input_path.with_stem(input_path.stem + "_output")
        )
        with open(input_path, newline="", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))

        results = []
        for i, row in enumerate(rows):
            mname = row.get("name", "").strip()
            formula = row.get("formula", "").strip()
            if not mname or not formula:
                results.append({**row, "json": "", "matched_template": "", "confidence": ""})
                continue
            print(f"  [{i+1}/{len(rows)}] {mname}", file=sys.stderr)
            try:
                steps, matched, score = convert(mname, formula, examples)
                results.append({
                    **row,
                    "json": json.dumps(steps),
                    "matched_template": matched["name"],
                    "confidence": f"{score:.0%}",
                })
            except Exception as e:
                print(f"    ERROR: {e}", file=sys.stderr)
                results.append({**row, "json": f"ERROR: {e}", "matched_template": "", "confidence": ""})

        fieldnames = list(rows[0].keys()) + ["json", "matched_template", "confidence"]
        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(results)
        print(f"\nOutput: {output_path}", file=sys.stderr)
        return

    # Single mode
    formula = args.formula or (sys.stdin.read().strip() if not sys.stdin.isatty() else input("Formula: ").strip())
    name = args.name or input("Measure name: ").strip()

    steps, matched, score = convert(name, formula, examples)
    print(f"Matched template: '{matched['name']}' (confidence: {score:.0%})", file=sys.stderr)
    print(json.dumps(steps, indent=2))


if __name__ == "__main__":
    main()
